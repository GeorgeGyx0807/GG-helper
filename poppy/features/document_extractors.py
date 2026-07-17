"""Local extractors for documents that are not plain text.

The personal library stores derived text, never a second copy of the user's
source file.  Each extractor returns bounded chunks with a small locator so
search results can say "page 3" or "Sheet1, rows 2-40" instead of pretending
every binary document has source-code line numbers.
"""

from dataclasses import dataclass
from functools import lru_cache
import json
import logging
import os
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path


DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".xls"}
MAX_EXTRACTED_CHARS = 2_000_000
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLUMNS = 200
ROWS_PER_CHUNK = 80


class DocumentExtractionError(ValueError):
    """A user-facing, safe-to-display extraction failure."""


@dataclass(frozen=True)
class ExtractedDocument:
    content: str
    chunks: list[dict]
    extractor: str


def extract_document(path, text_extensions):
    """Extract a supported path once and reuse it until the file changes."""

    path = Path(path).expanduser().resolve()
    try:
        stat = path.stat()
    except OSError as exc:
        raise DocumentExtractionError(f"无法读取文件: {exc}") from exc
    return _extract_document_cached(
        str(path),
        int(stat.st_size),
        int(stat.st_mtime_ns),
        tuple(sorted(str(item).lower() for item in text_extensions)),
    )


@lru_cache(maxsize=64)
def _extract_document_cached(path_text, _size, _mtime_ns, text_extensions):
    # Size and mtime are intentionally part of the cache key. A modified file
    # gets a fresh extraction while unchanged large documents become instant
    # on subsequent search/read calls in the same Gateway process.
    path = Path(path_text)
    extension = path.suffix.lower()
    if extension in text_extensions:
        return _extract_plain_text(path)
    if extension == ".pdf":
        return _extract_pdf(path)
    if extension in {".docx", ".pptx"}:
        return _extract_markitdown(path, extension)
    if extension in {".xlsx", ".xls"}:
        return _extract_spreadsheet(path, extension)
    raise DocumentExtractionError(f"不支持的文档格式: {extension or 'unknown'}")


def _extract_plain_text(path):
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise DocumentExtractionError(f"无法读取文件: {exc}") from exc
    content = _bounded(content)
    if not content.strip():
        raise DocumentExtractionError("文件没有可索引的文字")
    return ExtractedDocument(content, _cap_chunks(_line_chunks(content, {"kind": "text"})), "plain-text")


def _extract_pdf(path):
    """Use PDFium for the fast path, with pdfplumber as compatibility fallback."""
    pdfium_error = None
    try:
        import pypdfium2 as pdfium

        page_chunks = []
        page_texts = []
        document = pdfium.PdfDocument(str(path))
        try:
            for page_index in range(len(document)):
                page = document[page_index]
                text_page = None
                try:
                    text_page = page.get_textpage()
                    text = _bounded(text_page.get_text_range().replace("\r\n", "\n"))
                finally:
                    if text_page is not None:
                        text_page.close()
                    page.close()
                if not text:
                    continue
                page_number = page_index + 1
                page_texts.append(text)
                page_chunks.extend(_line_chunks(text, {"kind": "pdf_page", "page": page_number}))
        finally:
            document.close()
        if page_chunks:
            content = _bounded("\n\n".join(page_texts))
            return ExtractedDocument(content, _cap_chunks(page_chunks), "pypdfium2")
    except Exception as exc:  # fallback handles malformed or unusual PDFs
        pdfium_error = exc

    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover - dependency is packaged
        raise DocumentExtractionError("缺少 PDF 解析依赖，请重新安装应用") from (pdfium_error or exc)
    page_chunks = []
    page_texts = []
    pdfminer_logger = logging.getLogger("pdfminer")
    previous_log_level = pdfminer_logger.level
    pdfminer_logger.setLevel(logging.ERROR)
    try:
        try:
            with pdfplumber.open(str(path)) as document:
                for page_number, page in enumerate(document.pages, start=1):
                    text = _bounded(page.extract_text() or "")
                    if not text:
                        continue
                    page_texts.append(text)
                    page_chunks.extend(
                        _line_chunks(text, {"kind": "pdf_page", "page": page_number})
                    )
        finally:
            pdfminer_logger.setLevel(previous_log_level)
    except Exception as exc:
        raise DocumentExtractionError(f"PDF 解析失败: {exc}") from exc

    if not page_chunks:
        ocr_result = _extract_pdf_with_native_ocr(path)
        if ocr_result is not None:
            return ocr_result
        raise DocumentExtractionError("PDF 没有可提取的文字，且本机 OCR 不可用")
    content = _bounded("\n\n".join(page_texts))
    return ExtractedDocument(content, _cap_chunks(page_chunks), "pdfplumber")


def _extract_pdf_with_native_ocr(path):
    """Use macOS Vision only for PDFs that have no usable text layer."""
    helper = _native_ocr_helper()
    if helper is None:
        return None
    try:
        import pypdfium2 as pdfium

        document = pdfium.PdfDocument(str(path))
        with tempfile.TemporaryDirectory(prefix="poppy-ocr-") as directory:
            image_paths = []
            try:
                for page_index in range(len(document)):
                    page = document[page_index]
                    try:
                        image_path = Path(directory) / f"page-{page_index + 1:05d}.png"
                        page.render(scale=2).to_pil().save(image_path, format="PNG")
                        image_paths.append(image_path)
                    finally:
                        page.close()
            finally:
                document.close()
            completed = subprocess.run(
                [str(helper), *(str(item) for item in image_paths)],
                check=True,
                capture_output=True,
                text=True,
                timeout=300,
            )
        texts = json.loads(completed.stdout)
    except (OSError, subprocess.SubprocessError, ValueError, TypeError, json.JSONDecodeError):
        return None

    page_chunks = []
    page_texts = []
    for page_number, raw_text in enumerate(texts, start=1):
        text = _bounded(str(raw_text))
        if not text.strip():
            continue
        page_texts.append(text)
        page_chunks.extend(_line_chunks(text, {"kind": "pdf_page", "page": page_number}))
    if not page_chunks:
        return None
    return ExtractedDocument(
        _bounded("\n\n".join(page_texts)),
        _cap_chunks(page_chunks),
        "macos-vision-ocr",
    )


def _native_ocr_helper():
    configured = os.environ.get("POPPY_OCR_HELPER", "").strip()
    candidates = [Path(configured)] if configured else []
    bundle_root = getattr(sys, "_MEIPASS", "")
    if bundle_root:
        candidates.append(Path(bundle_root) / "poppy-ocr")
    for candidate in candidates:
        if candidate.is_file() and os.access(candidate, os.X_OK):
            return candidate
    return None


def _extract_markitdown(path, extension):
    try:
        from markitdown import MarkItDown
    except ImportError as exc:  # pragma: no cover - dependency is packaged
        label = "Word" if extension == ".docx" else "PowerPoint"
        raise DocumentExtractionError(f"缺少 {label} 解析依赖，请重新安装应用") from exc

    label = "Word" if extension == ".docx" else "PowerPoint"
    kind = "docx" if extension == ".docx" else "pptx"
    try:
        result = MarkItDown(enable_plugins=False).convert(str(path))
        content = _bounded(getattr(result, "markdown", getattr(result, "text_content", "")))
    except Exception as exc:
        raise DocumentExtractionError(f"{label} 文档解析失败: {exc}") from exc
    if not content.strip():
        raise DocumentExtractionError(f"{label} 文档没有可索引的文字")
    chunks = _cap_chunks(_line_chunks(content, {"kind": kind}))
    return ExtractedDocument(content, chunks, f"markitdown-{kind}")


def _extract_spreadsheet(path, extension):
    if extension == ".xlsx":
        sheets = _read_xlsx(path)
        extractor = "openpyxl"
    else:
        sheets = _read_xls(path)
        extractor = "xlrd"

    chunks = []
    sections = []
    for sheet_name, rows in sheets:
        if not rows:
            continue
        section_lines = [f"## 工作表: {sheet_name}"]
        section_lines.extend(f"第 {row_number} 行: " + " | ".join(values) for row_number, values in rows)
        section = "\n".join(section_lines)
        sections.append(section)
        for offset in range(0, len(rows), ROWS_PER_CHUNK):
            batch = rows[offset : offset + ROWS_PER_CHUNK]
            body = "\n".join(
                f"第 {row_number} 行: " + " | ".join(values)
                for row_number, values in batch
            ).strip()
            if body:
                chunks.append(
                    {
                        "line_start": int(batch[0][0]),
                        "line_end": int(batch[-1][0]),
                        "content": f"## 工作表: {sheet_name}\n{body}",
                        "location": {
                            "kind": "spreadsheet",
                            "sheet": str(sheet_name),
                            "row_start": int(batch[0][0]),
                            "row_end": int(batch[-1][0]),
                        },
                    }
                )

    content = _bounded("\n\n".join(sections))
    chunks = _cap_chunks(chunks)
    if not content.strip() or not chunks:
        raise DocumentExtractionError("表格没有可索引的单元格内容")
    return ExtractedDocument(content, chunks, extractor)


def _read_xlsx(path):
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency is packaged
        raise DocumentExtractionError("缺少 XLSX 解析依赖，请重新安装应用") from exc

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        output = []
        for sheet in workbook.worksheets:
            rows = []
            for row_number, row in enumerate(
                sheet.iter_rows(max_row=MAX_SHEET_ROWS, max_col=MAX_SHEET_COLUMNS, values_only=True),
                start=1,
            ):
                values = [_cell_text(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append((row_number, values))
            output.append((sheet.title, rows))
        workbook.close()
        return output
    except Exception as exc:
        raise DocumentExtractionError(f"Excel 文档解析失败: {exc}") from exc


def _read_xls(path):
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - dependency is packaged
        raise DocumentExtractionError("缺少 XLS 解析依赖，请重新安装应用") from exc

    try:
        workbook = xlrd.open_workbook(path, on_demand=True)
        output = []
        for sheet in workbook.sheets():
            rows = []
            for row_number in range(min(sheet.nrows, MAX_SHEET_ROWS)):
                values = [_cell_text(value) for value in sheet.row_values(row_number, 0, MAX_SHEET_COLUMNS)]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append((row_number + 1, values))
            output.append((sheet.name, rows))
        return output
    except Exception as exc:
        raise DocumentExtractionError(f"Excel 文档解析失败: {exc}") from exc


def _line_chunks(content, location, lines_per_chunk=ROWS_PER_CHUNK):
    lines = content.splitlines()
    chunks = []
    for start in range(0, len(lines), lines_per_chunk):
        body = "\n".join(lines[start : start + lines_per_chunk]).strip()
        if not body:
            continue
        line_start = start + 1
        line_end = min(len(lines), start + lines_per_chunk)
        chunk_location = dict(location)
        chunk_location.update({"line_start": line_start, "line_end": line_end})
        chunks.append(
            {
                "line_start": line_start,
                "line_end": line_end,
                "content": body,
                "location": chunk_location,
            }
        )
    return chunks


def _bounded(content):
    # PDF text extractors often return CJK compatibility radicals and NUL
    # separators instead of the displayed characters. NFKC fixes the common
    # compatibility forms (e.g. ⼀ -> 一) and removing NULs keeps search usable.
    normalized = unicodedata.normalize("NFKC", str(content or ""))
    normalized = normalized.replace("\x00", "")
    return normalized[:MAX_EXTRACTED_CHARS].strip()


def _cap_chunks(chunks):
    """Keep FTS chunks within the same extraction budget as document text."""

    remaining = MAX_EXTRACTED_CHARS
    bounded = []
    for item in chunks:
        if remaining <= 0:
            break
        content = str(item.get("content") or "")[:remaining].strip()
        if not content:
            continue
        bounded.append({**item, "content": content})
        remaining -= len(content)
    return bounded


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
